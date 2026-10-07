"""Prepare a source-locked BMC Research Note; never submit or invent declarations."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import re
import tempfile
from importlib.metadata import version
from pathlib import Path
from typing import Any

import build_path_b_documents as path_b
import build_research_documents as documents
import numpy as np
from author_release import CONFIRMATIONS as CONFIRMATIONS
from author_release import STATEMENTS as STATEMENTS
from author_release import blank_record, unique_object, valid_statement
from author_release import release_gaps as release_gaps
from bmc_prose import validate_prose
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from selectivity import json_bytes, publish
from submission_archive import code_access, distribution_gaps, source_bytes
from submission_archive import source_archives as source_archives
from submission_archive import zip_payload as zip_payload
from threadpoolctl import threadpool_info

PACKAGE = Path("research/submission")
MAX_ADDITIONAL_BYTES = 20_000_000


def article_checks(text: str) -> dict[str, Any]:
    headings = (
        "Abstract",
        "Keywords",
        "Introduction",
        "Main text",
        "Limitations",
        "List of abbreviations",
        "References",
    )
    positions = []
    for heading in headings:
        marker = f"## {heading}\n"
        if text.count(marker) != 1:
            raise ValueError(f"Missing or duplicate journal section: {heading}")
        positions.append(text.index(marker))
    if positions != sorted(positions):
        raise ValueError("Journal sections are out of order")
    abstract = text.split("## Abstract\n")[1].split("## Keywords\n")[0]
    if any(abstract.count(f"### {h}\n") != 1 for h in ("Objective", "Results")):
        raise ValueError("Abstract requires Objective and Results")
    if re.search(r"\[\d", abstract):
        raise ValueError("Abstract must not contain reference citations")
    body = text.split("## Introduction\n")[1].split("## List of abbreviations\n")[0]
    prose = "\n".join(line for line in body.splitlines() if not line.startswith(("|", "<!--")))
    abstract_words, body_words = (
        len(re.sub(r"(?m)^#{1,6}\s+", "", documents.plain(abstract)).split()),
        len(re.sub(r"(?m)^#{1,6}\s+", "", documents.plain(prose)).split()),
    )
    keywords = text.split("## Keywords\n")[1].split("## Introduction\n")[0].strip().split(";")
    tables = [b for b in documents.read_blocks(text) if b[0] == "table"]
    if abstract_words > 200 or body_words > 2000 or not 3 <= len(keywords) <= 10:
        raise ValueError("Journal word or keyword limit exceeded")
    if len(tables) != 3 or re.search(r"^Figure \d", text, re.MULTILINE):
        raise ValueError("This note requires exactly three tables and no main figures")
    citations = re.findall(r"Table ([1-3])(?!\.)", body)
    if list(dict.fromkeys(citations)) != ["1", "2", "3"]:
        raise ValueError("Tables must be cited in order")
    for number in range(1, 4):
        match = re.search(
            rf"^Table {number}\. ([^\n]+)\n\n<!-- path-b:[^\n]+:start -->"
            r"\n.*?\n<!-- path-b:[^\n]+:end -->\n\n([^\n]+)",
            text,
            re.S | re.M,
        )
        if not match or len(match[1].split()) > 15 or len(match[2].split()) > 300:
            raise ValueError("Invalid table title or legend")
    return {
        "abstract_words_including_subheadings": abstract_words,
        "body_words_including_headings_titles_legends_excluding_cells": body_words,
        "keywords": len(keywords),
        "main_tables": len(tables),
        "main_figures": 0,
    }


def format_journal_docx(path: Path) -> None:
    doc = Document(str(path))
    doc.styles["Normal"].font.size = Pt(12)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(16)
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.line_spacing = 2
        paragraph.paragraph_format.page_break_before = False
    for table in doc.tables:
        table.style = "Table Grid"
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 2
                    paragraph.paragraph_format.keep_with_next = False
    for section in doc.sections:
        numbering = OxmlElement("w:lnNumType")
        for key, value in {"countBy": "1", "start": "0", "restart": "continuous"}.items():
            numbering.set(qn(f"w:{key}"), value)
        section._sectPr.append(numbering)
        footer = section.footer.paragraphs[0]
        footer.text = "Page "
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    doc.save(str(path))


def declarations(
    record: dict[str, Any], archive_hashes: dict[str, str], access: dict[str, Any] | None = None
) -> str:
    statement = record["statements"]
    pairs = (
        ("Ethics approval and consent to participate", statement["ethics_approval_and_consent"]),
        ("Consent for publication", statement["consent_for_publication"]),
        (
            "Availability of data and materials",
            "The source snapshots, code and recorded results "
            "supporting this reanalysis accompany Additional files 2 and 3. The historical "
            "18,412-compound library and raw catalytic replicate vectors are unavailable. "
            + " ".join(f"{name} SHA-256: {digest}." for name, digest in archive_hashes.items())
            + (
                " Code access and requirements: " + json.dumps(access, ensure_ascii=False)
                if access
                else ""
            )
            + " Rights/license statement: "
            + str(statement["rights_permissions_and_code_license"]),
        ),
        ("Competing interests", statement["competing_interests"]),
        ("Funding", statement["funding"]),
        ("Authors' contributions", statement["author_contributions"]),
        ("Acknowledgements", statement["acknowledgements"]),
        ("AI assistance", statement["ai_assistance_disclosure"]),
    )
    return "## Declarations\n\n" + "\n\n".join(f"### {h}\n\n{v}" for h, v in pairs) + "\n\n"


def build(
    repository: Path,
    output: Path,
    *,
    require_author_confirmation: bool = False,
    author_record: Path | None = None,
) -> Path:
    if output.is_symlink() or (output.exists() and (not output.is_dir() or any(output.iterdir()))):
        raise ValueError("Submission output must be a new or empty directory")
    path_b.locked_inputs(repository, repository / "research/results")
    package = repository / PACKAGE
    text = (package / "manuscript.md").read_text()
    inputs = [
        repository / name
        for name in (
            "research/structure_comparison/results/observed_proximity.json",
            "research/results/selectivity.json",
            "research/structure_comparison/model_support/results/model_support.json",
            "research/results/controls.json",
        )
    ]
    data = [json.loads(p.read_text()) for p in inputs]
    validate_prose(repository, text, data)
    blocks = path_b.quantitative_blocks(*data)
    path_b.synchronize(text, {k: blocks[k] for k in ("sites", "paired", "controls")})
    checks = article_checks(text)
    record_path = (
        author_record if author_record is not None else package / "author_confirmation.json"
    )
    record_bytes = (
        source_bytes(record_path.parent, record_path.name)
        if author_record is not None or record_path.exists()
        else json_bytes(blank_record())
    )
    record = json.loads(record_bytes, object_pairs_hook=unique_object)
    gaps = release_gaps(record)
    rights_gaps = distribution_gaps(repository)
    access = code_access(repository)
    if require_author_confirmation and (gaps or rights_gaps):
        raise ValueError("Author release blocked: " + ", ".join(gaps + rights_gaps))
    cover_path = package / "cover_letter.md"
    if not cover_path.exists():
        cover_path = package / "cover_letter_core.md"
    cover = cover_path.read_text()
    if len(cover.split("Dear Editors,\n", 1)[1].split()) != 250:
        raise ValueError("Cover-letter scientific core must be 250 words")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".bmc-note-", dir=output.parent) as folder:
        staging = Path(folder)
        private_source = repository
        if not (repository / ".git").exists() and (repository / "SHA256SUMS.json").exists():
            private_source = repository.parent / "private/source"
        private_manifest = json.loads((package / "private_archive.json").read_text())
        private_files = {}
        for name, expected in private_manifest["files"].items():
            payload = source_bytes(private_source, name)
            if hashlib.sha256(payload).hexdigest() != expected:
                raise ValueError(f"Private provenance changed: {name}")
            private_files[f"private/source/{name}"] = payload
        (staging / "Private_audit_provenance.zip").write_bytes(zip_payload(private_files))
        baseline = staging / "path_b"
        documents.build(
            repository / "research/manuscript.md",
            repository / "research/results",
            baseline,
            profile="path-b",
            repository=repository,
            structure_input=inputs[0],
            structure_manifest=inputs[0].parent / "derived/derived_manifest.json",
            handoff_directory=repository / "research/structure_comparison",
        )
        # The supplement keeps its own original numbered references and scientific limitations.
        original = (repository / "research/manuscript.md").read_text()
        methods = re.search(r"^## 2\..*?\n(.*?)^## 3\.", original, re.M | re.S)
        if methods is None:
            raise ValueError("Cannot locate source-locked structural methods")
        supplement = (baseline / "diagnostic_supplement.md").read_text()
        supplement = re.sub(
            r"^Author affiliation, funding, COI,.*?\n\n", "", supplement, flags=re.M | re.S
        )
        supplement = (
            "# Additional file 1: Supplementary methods and diagnostic results\n\n"
            "This supplement preserves scientific limitations and source-reported uncertainty. "
            "Author declarations and approvals are handled in the main-document release.\n\n"
            "## Extended structural and endpoint methods\n\n"
            + methods[1]
            + "\n\n"
            + supplement
            + "\n\n## Supplementary references\n\n"
            + original.split("## References\n", 1)[1]
        )
        (staging / "Additional_file_1.md").write_text(supplement)
        figures = []
        for heading, path, caption in _supplement_figures(baseline):
            figures.append((heading, path, caption))
        documents.render_document(
            documents.read_blocks(supplement),
            staging,
            figures,
            path_b_layout=True,
            stem="Additional_file_1",
            title="Supplementary methods and diagnostics",
        )
        first, second = source_archives(repository, baseline)
        archives = {"Additional_file_2.zip": first, "Additional_file_3.zip": second}
        for name, payload in archives.items():
            (staging / name).write_bytes(payload)
        if not gaps and not rights_gaps:
            text = re.sub(
                r"^AUTHOR-REVIEW DRAFT\..*?$",
                lambda _: record["statements"]["final_byline_and_addresses"],
                text,
                flags=re.M,
            )
            text = text.replace(
                "## References\n",
                declarations(
                    record,
                    {
                        name: hashlib.sha256(payload).hexdigest()
                        for name, payload in archives.items()
                    },
                    access,
                )
                + "## References\n",
            )
            cover = "Dear Editors,\n" + cover.split("Dear Editors,\n", 1)[1]
            cover += "\n\nCompeting interests: " + record["statements"]["competing_interests"]
            cover += (
                "\n\nPrior versions and submissions: "
                + record["statements"]["prior_publication_and_submission_history"]
            )
            cover += (
                "\n\nAll authors have approved this manuscript and accept accountability. "
                "No concurrent submission is active.\n"
            )
        else:
            byline = record["statements"]["final_byline_and_addresses"]
            if valid_statement(byline):
                text = re.sub(
                    r"^AUTHOR-REVIEW DRAFT\..*?$",
                    lambda match: byline + "\n\n" + match[0],
                    text,
                    flags=re.M,
                )
            cover = (
                "# Cover-letter scientific core\n\n"
                "AUTHOR-REVIEW DRAFT. The text below is not a complete submission letter. "
                "Required statements about competing interests, author approval and "
                "publication/submission history must come from the author. "
                "No statement of exclusivity or approval is implied.\n\n"
                "Dear Editors,\n" + cover.split("Dear Editors,\n", 1)[1]
            )
        checks = article_checks(text)
        (staging / "BMC_research_note.md").write_text(text)
        (staging / "Cover_letter.md").write_text(cover)
        for stem, content in (("BMC_research_note", text), ("Cover_letter", cover)):
            _, word = documents.render_document(
                documents.read_blocks(content),
                staging,
                [],
                path_b_layout=True,
                stem=stem,
                title="BMC Research Notes preparation",
                literal_paragraphs=frozenset(
                    b[1]
                    for b in documents.read_blocks(content)
                    if b[0] == "paragraph"
                    and any(
                        isinstance(value, str) and value in b[1]
                        for value in record["statements"].values()
                    )
                ),
                table_width_overrides={
                    "Control / method": [155, 105, 110, 110],
                    "PDB": [40, 45, 115, 44, 44, 48, 74, 70],
                },
            )
            format_journal_docx(word)
        for name in (
            "checklist.md",
            "author_actions.md",
            "reviewer_candidates.md",
            "claim_traceability.md",
        ):
            if (package / name).exists():
                (staging / name).write_bytes((package / name).read_bytes())
        (staging / "author_confirmation.json").write_bytes(record_bytes)
        for name in ("Additional_file_1.pdf", *archives):
            if (staging / name).stat().st_size > MAX_ADDITIONAL_BYTES:
                raise ValueError(f"Additional file exceeds 20 MB: {name}")
        payloads = {p.name: p.read_bytes() for p in staging.iterdir() if p.is_file()}
        manifest = {
            "journal": "BMC Research Notes",
            "article_type": "Research note",
            "author_release_checks_complete": not gaps,
            "author_record_sha256": hashlib.sha256(record_bytes).hexdigest(),
            "external_author_record": author_record is not None,
            "submission_ready": not gaps and not rights_gaps,
            "distribution_gaps": rights_gaps,
            "code_access": access,
            "runtime": {
                "python": platform.python_version(),
                "platform": platform.platform(),
                "versions": {
                    name: version(name)
                    for name in ("numpy", "scipy", "scikit-learn", "rdkit", "gemmi")
                },
                "numpy_build": np.show_config(mode="dicts"),
                "threadpools": [
                    {k: v for k, v in pool.items() if k != "filepath"} for pool in threadpool_info()
                ],
                "thread_environment": {
                    name: os.environ.get(name)
                    for name in (
                        "OMP_NUM_THREADS",
                        "OPENBLAS_NUM_THREADS",
                        "MKL_NUM_THREADS",
                        "NPY_DISABLE_CPU_FEATURES",
                    )
                },
                "scope": (
                    "Current rendering/check environment, not historical model-fit provenance "
                    "or cross-platform equivalence"
                ),
            },
            "submitted": False,
            "biological_validation": False,
            "author_owned_gaps": gaps,
            "qualification": (
                "Checks verify documented confirmations, not their truth or editorial acceptance."
            ),
            "article_checks": checks,
            "cover_letter_scientific_core_words": 250,
            "source_lock_sha256": path_b.LOCK_SHA256,
            "files": {
                name: {"sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)}
                for name, data in sorted(payloads.items())
            },
        }
        payloads["submission-manifest.json"] = json_bytes(manifest)
        publish(output, payloads, "submission-manifest.json")
    return output / "submission-manifest.json"


def _supplement_figures(baseline: Path) -> list[tuple[str, Path, str]]:
    """Recover exact captions and image relationships from the generated supplement DOCX."""
    doc = Document(str(baseline / "Path_B_diagnostic_supplement.docx"))
    figures: list[tuple[str, Path, str]] = []
    for index, paragraph in enumerate(doc.paragraphs):
        drawings = paragraph._p.xpath(".//w:drawing")
        if not drawings:
            continue
        if len(drawings) != 1 or not 0 < index < len(doc.paragraphs) - 1:
            raise ValueError("Supplement figure/caption mapping is incomplete")
        heading, caption = doc.paragraphs[index - 1].text, doc.paragraphs[index + 1].text
        if not heading.startswith("Figure S") or not caption:
            raise ValueError("Supplement figure/caption mapping is incomplete")
        blips = paragraph._p.xpath(".//a:blip")
        if len(blips) != 1:
            raise ValueError("Ambiguous supplement image")
        relationship = blips[0].get(qn("r:embed"))
        if relationship is None:
            raise ValueError("Missing supplement image relationship")
        part = doc.part.related_parts[relationship]
        path = baseline / f"supplement-image-{len(figures)}.png"
        path.write_bytes(part.blob)
        figures.append((heading, path, caption))
    if not figures or len(figures) != len(doc.inline_shapes):
        raise ValueError("Supplement figure/caption mapping is incomplete")
    return figures


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repository", type=Path, default=documents.ROOT)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--require-author-confirmation", action="store_true")
    parser.add_argument(
        "--author-record",
        type=Path,
        help="Completed author JSON outside the frozen provenance; never inferred as approval",
    )
    args = parser.parse_args()
    build(
        args.repository,
        args.output,
        require_author_confirmation=args.require_author_confirmation,
        author_record=args.author_record,
    )


if __name__ == "__main__":
    main()
