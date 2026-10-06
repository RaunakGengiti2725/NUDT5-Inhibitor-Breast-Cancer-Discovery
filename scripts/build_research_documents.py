"""Render the authored Markdown revision and recorded diagnostics into PDF/DOCX."""

from __future__ import annotations

import argparse
import html
import importlib
import json
import re
import tempfile
from pathlib import Path
from typing import Any

import build_structure_comparison_figures as structure_figures
from build_selectivity_figures import FIGURE_NAMES, load_recorded, render
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from selectivity import artifacts, json_bytes, publish, run_manifest

ROOT = Path(__file__).resolve().parents[1]
STRUCTURE_CAPTION = (
    "Observed-coordinate comparison of published compound-9 complexes (Balıkçı et al. 2024, "
    "Figures 3–4; PDB 8RIY/8OTV). Each target-specific row has a retained pair within 5.0 Å "
    "at either site; columns preserve both sites and both dimer chains. Inclusive radii are "
    "3.5/4.0/4.5/5.0 Å (4.0 primary), with no threshold tuning. Grey cells exceed 5.0 Å; "
    "null/refused rows are counted below, not plotted as no-contact. Partial residues (†) give "
    "upper bounds on unknown complete-residue minima; positive fractional occupancies (*) "
    "are unweighted. Residue axes are not homology-aligned; crystal copies are not independent n. "
    "No density or coordinate-uncertainty propagation. Residue-name proximity is not a "
    "guanidinium interaction, binding energy, hydrogen bond, selectivity or a causal test of "
    "the published Arg51 rationale. Full-precision tables retain every row, including nulls."
)


def read_blocks(text: str) -> list[tuple[str, Any]]:
    lines = text.splitlines()
    blocks: list[tuple[str, Any]] = []
    paragraph: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("<!--") and line.endswith("-->"):
            index += 1
            continue
        if not line:
            if paragraph:
                blocks.append(("paragraph", " ".join(paragraph)))
                paragraph = []
        elif line.startswith("#") or line.startswith("|"):
            if paragraph:
                blocks.append(("paragraph", " ".join(paragraph)))
                paragraph = []
            if line.startswith("#"):
                blocks.append(
                    ("heading", (len(line) - len(line.lstrip("#")), line.lstrip("#").strip()))
                )
            else:
                rows = []
                while index < len(lines) and lines[index].strip().startswith("|"):
                    cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                        rows.append(cells)
                    index += 1
                blocks.append(("table", rows))
                continue
        else:
            paragraph.append(line)
        index += 1
    if paragraph:
        blocks.append(("paragraph", " ".join(paragraph)))
    return blocks


def plain(text: str) -> str:
    return re.sub(r"(?<!\w)\*([^*\n]+)\*(?!\w)", r"\1", text.replace("**", "").replace("`", ""))


def diagnostic_figure(results: Path, output: Path) -> Path:
    mpl = importlib.import_module("matplotlib")
    mpl.use("Agg")
    plt = importlib.import_module("matplotlib.pyplot")
    data = json.loads((results / "benchmark.json").read_text())
    methods = ["Property_LR", "SVM_RBF", "RF", "Equal_mean", "Nearest_active", "GBT"]
    labels = ["Property LR", "RBF-SVM", "RF", "Equal mean", "Nearest active", "GBT"]
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.7), layout="constrained")
    for index, split in enumerate(["molecule", "scaffold", "series"]):
        points = [i + (index - 1) * 0.24 for i in range(len(methods))]
        axes[0].bar(
            points,
            [data[split]["metrics"][method]["auc"] for method in methods],
            width=0.24,
            label=split,
            color=["#b2becd", "#527893", "#bf755e"][index],
        )
    axes[0].set_xticks(range(len(labels)), labels, rotation=40, ha="right")
    axes[0].set_ylim(0, 1.05)
    axes[0].axhline(0.5, color="#888888", linestyle=":", linewidth=1)
    axes[0].set_ylabel("Pooled out-of-fold ROC-AUC")
    axes[0].set_title("A  Seed 42; small, unverified-label dataset", fontsize=10)
    axes[0].legend(frameon=False, fontsize=8)
    for method, colour in [("Property_LR", "#527893"), ("Equal_mean", "#bf755e")]:
        seed_rows = [row for row in data["seed_sensitivity"] if row["split"] == "series"]
        axes[1].plot(
            [row["seed"] for row in seed_rows],
            [row["metrics"][method]["auc"] for row in seed_rows],
            marker="o",
            color=colour,
            label=method,
        )
    axes[1].set_ylim(0, 1.05)
    axes[1].axhline(0.5, color="#888888", linestyle=":", linewidth=1)
    axes[1].set_xlabel("Seed; no best-seed selection")
    axes[1].set_ylabel("Pooled series-holdout ROC-AUC")
    axes[1].set_title("B  Sensitivity to decoy partitions", fontsize=10)
    axes[1].set_xticks(sorted({row["seed"] for row in data["seed_sensitivity"]}))
    axes[1].legend(frameon=False, fontsize=8, loc="center right")
    for ax in axes:
        ax.spines[["top", "right"]].set_visible(False)
    path = output / "diagnostics.png"
    fig.savefig(path, dpi=180)
    plt.close(fig)
    return path


def _build(
    manuscript: Path,
    results: Path,
    output: Path,
    paired: dict[str, Any] | None,
    structural: dict[str, Any] | None,
    handoff: dict[str, bytes],
) -> tuple[Path, Path]:
    if output.exists() and (not output.is_dir() or any(output.iterdir())):
        raise ValueError("Document output must be a new or empty directory")
    blocks = read_blocks(manuscript.read_text())
    if not blocks:
        raise ValueError("Manuscript is empty")
    if not (results / "benchmark.json").is_file():
        raise ValueError("Recorded benchmark.json is required; results are never fabricated")
    if not (results / "controls.json").is_file() or not (results / "transfer.json").is_file():
        raise ValueError("Recorded controls.json and transfer.json are required")
    output.mkdir(parents=True, exist_ok=True)
    figure = diagnostic_figure(results, output)
    extension_figures = importlib.import_module("build_extension_figures")
    extension = extension_figures.build_figures(results, output)
    extension_figures.write_tables(results, output)
    if paired is not None:
        for name, content in {**render(paired), **artifacts(paired)}.items():
            (output / name).write_bytes(content)
        for scenario, name in zip(paired["summary"]["scenarios"], FIGURE_NAMES, strict=True):
            n = paired["summary"]["scenarios"][scenario]["n"]
            extension.append(
                (
                    output / f"{name}.png",
                    f"Retrospective paired-target evidence ({scenario}; n={n} nonoverlap pairs). "
                    "R = reported mean catalytic IC50(NUDT14)/IC50(NUDT5), dimensionless; "
                    "target IC50 values in µM and source SDs are retained in pharmacology.csv. "
                    "Arrows are strict censoring bounds, not confidence intervals; crosses have "
                    "no finite ratio bound. Frozen Equal_mean is an uncalibrated label score. "
                    "Reaction times and control/replication ambiguities limit comparison. "
                    + ("No eligible pairs. " if not n else "")
                    + paired["warning"],
                )
            )
    else:
        blocks.append(
            (
                "paragraph",
                "Paired-target results unavailable: no selectivity figure or table was generated. "
                "Unavailable measurements are not zero; legacy diagnostic results remain separate.",
            )
        )
    structural_headings: dict[Path, str] = {}
    if structural is not None:
        for name, content in {
            **structure_figures.tables(structural),
            **structure_figures.render(structural),
            **handoff,
        }.items():
            (output / name).write_bytes(content)
        figure_number = len(extension) + 2
        for index, target in enumerate(sorted(s["target"] for s in structural["structures"])):
            for name, content in structure_figures.render(structural, target=target).items():
                (output / name).write_bytes(content)
            path = output / f"{structure_figures.FIGURE}_{target}.png"
            structural_headings[path] = f"Figure {figure_number}{chr(65 + index)}"
            extension.append(
                (
                    path,
                    f"{target}: published compound-9 coordinates [2], not new "
                    "binding measurements. Both sites and every row with any retained "
                    "pair within 5 Å remain. Missingness and limitations are adjacent; "
                    "full vector map and all-residue tables are in the supplement.",
                )
            )
        with (output / "supplementary_results.md").open("a") as stream:
            stream.write(
                "\n\n## Observed-coordinate supplement\n\n"
                + STRUCTURE_CAPTION
                + "\n\nFull residue_proximity.csv, atom_pairs_within_5A.csv and "
                "radius_sensitivity.csv retain all sites and missingness. "
                "lab_handoff.md, hypotheses_controls.csv and handoff_sources.json "
                "separate proposals, source Methods and unresolved prerequisites.\n"
            )
    else:
        blocks.append(
            (
                "paragraph",
                "Structural results unavailable: no structural figure, "
                "distance table or lab handoff was generated. Missing evidence is not "
                "zero proximity or no contact; no geometry claim is supplied by this build.",
            )
        )
    caption = (
        "Figure 1. Newly computed diagnostics, not historical or prospective validation. "
        "A: seed-42 pooled scores; series pooling mixes differently trained models. "
        "B: all five specified seeds. No interval denotes population uncertainty. "
        "Descriptor performance is compatible with confounding but does not establish "
        "its causal contribution."
    )
    figures = [("Diagnostic figure", figure, caption)]
    figures.extend(
        (structural_headings.get(path, f"Figure {index}"), path, caption)
        for index, (path, caption) in enumerate(extension, start=2)
    )
    return render_document(blocks, output, figures)


def render_document(
    blocks: list[tuple[str, Any]],
    output: Path,
    figures: list[tuple[str, Path, str]],
    *,
    path_b_layout: bool = False,
    stem: str = "NUDT5_evidence_bounded_revision",
    title: str = "NUDT5 reproducibility and chemical-identity controls",
) -> tuple[Path, Path]:
    """Render explicit content/figure order; no implicit diagnostic appendices."""
    flow = importlib.import_module("reportlab.platypus")
    styles_module = importlib.import_module("reportlab.lib.styles")
    colours = importlib.import_module("reportlab.lib.colors")
    pdfmetrics = importlib.import_module("reportlab.pdfbase.pdfmetrics")
    fonts = importlib.import_module("reportlab.pdfbase.ttfonts")
    font_manager = importlib.import_module("matplotlib.font_manager")
    pdfmetrics.registerFont(fonts.TTFont("Research", font_manager.findfont("DejaVu Sans")))
    pdfmetrics.registerFont(
        fonts.TTFont(
            "ResearchBold",
            font_manager.findfont(font_manager.FontProperties(family="DejaVu Sans", weight="bold")),
        )
    )
    styles = styles_module.getSampleStyleSheet()
    for name in ["BodyText", "Normal", "Title", "Heading1", "Heading2", "Heading3"]:
        styles[name].fontName = (
            "ResearchBold" if name.startswith("Heading") or name == "Title" else "Research"
        )
        styles[name].fontSize = 9 if name in ["BodyText", "Normal"] else styles[name].fontSize
        styles[name].leading = 13 if name in ["BodyText", "Normal"] else styles[name].leading
        styles[name].spaceAfter = 7
    table_style = styles_module.ParagraphStyle(
        "TableBody",
        parent=styles["BodyText"],
        fontSize=8 if path_b_layout else 7,
        leading=11 if path_b_layout else 10,
    )
    for name in ("Title", "Heading1", "Heading2", "Heading3"):
        styles[name].keepWithNext = True
    doc = Document()
    doc.sections[0].left_margin = Inches(1)
    doc.sections[0].right_margin = Inches(1)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    if path_b_layout:
        section = doc.sections[0]
        section.left_margin = section.right_margin = Inches(0.75)
        section.top_margin = section.bottom_margin = Inches(0.75)
    story = []
    for kind, value in blocks:
        if kind == "heading":
            level, text = value
            doc.add_heading(plain(text), level=min(level - 1, 3))
            story.append(
                flow.Paragraph(
                    html.escape(plain(text)),
                    styles["Title" if level == 1 else f"Heading{min(level - 1, 3)}"],
                )
            )
        elif kind == "paragraph":
            doc.add_paragraph(plain(value))
            story.append(flow.Paragraph(html.escape(plain(value)), styles["BodyText"]))
        elif kind == "table":
            table = doc.add_table(rows=1, cols=len(value[0]))
            table.style = "Light Shading Accent 1"
            for index, cell in enumerate(value[0]):
                table.rows[0].cells[index].text = plain(cell)
            for row in value[1:]:
                for cell, text in zip(table.add_row().cells, row, strict=True):
                    cell.text = plain(text)
            widths = [168.0] + [312 / (len(value[0]) - 1)] * (len(value[0]) - 1)
            if path_b_layout:
                column_widths: dict[int, list[float]] = {
                    2: [300, 180],
                    4: [60, 140, 140, 140],
                    5: [145, 65, 65, 95, 110],
                    6: [35, 45, 180, 65, 65, 90],
                    7: [160, 100, 30, 48, 48, 47, 47],
                }
                widths = column_widths.get(len(value[0]), [480 / len(value[0])] * len(value[0]))
                table.style = "Table Grid"
                table.autofit = False
                table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                for column, width in zip(table.columns, widths, strict=True):
                    column.width = Pt(width)
                for row_index, word_row in enumerate(table.rows):
                    word_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                    for word_cell, width in zip(word_row.cells, widths, strict=True):
                        word_cell.width = Pt(width)
                        for paragraph in word_cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(2)
                            paragraph.paragraph_format.keep_with_next = (
                                len(value) < 25 and row_index < len(value) - 1
                            )
                            for run in paragraph.runs:
                                run.font.size = Pt(8)
                                run.bold = row_index == 0
            pdf_rows = [
                [flow.Paragraph(html.escape(plain(cell)), table_style) for cell in row]
                for row in value
            ]
            pdf_table = flow.Table(
                pdf_rows,
                colWidths=widths,
                repeatRows=1,
            )
            pdf_table.setStyle(
                flow.TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colours.HexColor("#e7edf2")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                        ("LINEBELOW", (0, 0), (-1, 0), 0.5, colours.HexColor("#8798aa")),
                    ]
                )
            )
            story.extend([pdf_table, flow.Spacer(1, 10)])
    for heading, path, figure_caption in figures:
        doc.add_heading(heading, level=1).paragraph_format.page_break_before = True
        image = flow.Image(str(path))
        scale = (
            min(480 / image.drawWidth, 500 / image.drawHeight)
            if path_b_layout
            else (490 / image.drawWidth)
        )
        image.drawWidth *= scale
        image.drawHeight *= scale
        if not path_b_layout and heading == "Diagnostic figure":
            image.drawHeight = 210
        if path_b_layout:
            doc.add_picture(str(path), width=Pt(image.drawWidth), height=Pt(image.drawHeight))
            picture_format = doc.paragraphs[-1].paragraph_format
            picture_format.keep_with_next = True
            picture_format.line_spacing = 1
            picture_format.space_after = Pt(6)
        else:
            doc.add_picture(str(path), width=Inches(6.4))
        caption_paragraph = doc.add_paragraph(figure_caption)
        if path_b_layout:
            caption_paragraph.paragraph_format.keep_together = True
            caption_paragraph.paragraph_format.line_spacing = 1
            for run in caption_paragraph.runs:
                run.font.size = Pt(9)
        story.extend(
            [
                flow.PageBreak(),
                flow.Paragraph(heading, styles["Heading1"]),
                image,
                flow.Paragraph(figure_caption, styles["BodyText"]),
            ]
        )
    pdf = output / f"{stem}.pdf"
    word = output / f"{stem}.docx"
    doc.save(str(word))
    template = flow.SimpleDocTemplate(
        str(pdf),
        pagesize=(612, 792),
        leftMargin=54,
        rightMargin=54,
        topMargin=48,
        bottomMargin=48,
        title=title,
    )

    def footer(canvas: Any, document: Any) -> None:
        canvas.setFont("Research", 7)
        canvas.setFillColor(colours.HexColor("#64748b"))
        canvas.drawString(54, 28, "Author-review revision | No new biological experiments")
        canvas.drawRightString(558, 28, str(document.page))

    template.build(story, onFirstPage=footer, onLaterPages=footer)
    return pdf, word


def build(
    manuscript: Path,
    results: Path,
    output: Path,
    *,
    require_selectivity: bool = False,
    structure_input: Path | None = None,
    structure_manifest: Path | None = None,
    require_structure: bool = False,
    repository: Path = ROOT,
    handoff_directory: Path = ROOT / "research/structure_comparison",
    profile: str = "legacy",
) -> tuple[Path, Path]:
    """Stage complete documents before non-overwriting, completion-last publication."""
    if output.is_symlink() or (output.exists() and (not output.is_dir() or any(output.iterdir()))):
        raise ValueError("Document output must be a new or empty directory, not a symlink")
    if not read_blocks(manuscript.read_text()):
        raise ValueError("Manuscript is empty")
    if profile not in {"legacy", "path-b"}:
        raise ValueError("Unknown document profile")
    if profile == "path-b":
        require_structure = require_selectivity = True
    inputs = [manuscript]
    for name in ("benchmark.json", "controls.json", "transfer.json"):
        path = results / name
        if not path.is_file():
            raise ValueError(f"Recorded {name} is required; results are never fabricated")
        inputs.append(path)
    source, manifest = results / "selectivity.json", results / "selectivity-manifest.json"
    paired = None
    if source.exists() or manifest.exists() or require_selectivity:
        if not source.is_file() or not manifest.is_file():
            raise ValueError("Recorded selectivity.json and selectivity-manifest.json are required")
        paired = load_recorded(source, manifest)
        inputs.extend([source, manifest])
    structural = None
    handoff: dict[str, bytes] = {}
    supplied = [path for path in (structure_input, structure_manifest) if path is not None]
    if require_structure or any(path.exists() or path.is_symlink() for path in supplied):
        if structure_input is None or structure_manifest is None:
            raise ValueError("Required structural input and manifest")
        structural = structure_figures.load_recorded(
            structure_input, structure_manifest, repository=repository
        )
        inputs.extend([structure_input, structure_manifest])
        inputs.extend(structure_figures.publication_inputs(repository))
        inputs.extend(
            repository / name for name in structure_figures.source_hashes(structural["provenance"])
        )
        for name in ("lab_handoff.md", "hypotheses_controls.csv", "handoff_sources.json"):
            path = handoff_directory / name
            if not path.is_file() or path.is_symlink() or not path.read_bytes().strip():
                raise ValueError(f"Required nonempty structural handoff: {path}")
            handoff[name] = path.read_bytes()
            inputs.append(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".nudt5-documents-", dir=output.parent) as directory:
        staging = Path(directory)
        if profile == "path-b":
            path_b = importlib.import_module("build_path_b_documents")
            pdf, word, extra_inputs = path_b.build(
                manuscript, results, staging, paired, structural, repository
            )
            inputs.extend(extra_inputs)
        else:
            pdf, word = _build(manuscript, results, staging, paired, structural, handoff)
        payloads = {path.name: path.read_bytes() for path in staging.iterdir()}
        inputs.extend(
            [
                Path(__file__),
                Path(__file__).with_name("build_extension_figures.py"),
                Path(__file__).with_name("build_selectivity_figures.py"),
                Path(__file__).with_name("build_structure_comparison_figures.py"),
                Path(__file__).parent / "scripts/structure_comparison.py",
                Path(__file__).parent / "scripts/selectivity.py",
                Path(__file__).parent / "scripts/pipeline.py",
            ]
        )
        record = run_manifest(
            inputs,
            {
                "require_selectivity": require_selectivity,
                "require_structure": require_structure,
                "profile": profile,
            },
            payloads,
        )
        record["structure_status"] = "recorded" if structural is not None else "unavailable"
        record["selectivity_status"] = "recorded" if paired is not None else "unavailable"
        payloads["documents-manifest.json"] = json_bytes(record)
        publish(output, payloads, "documents-manifest.json")
        return output / pdf.name, output / word.name


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--profile", choices=("legacy", "path-b"), default="legacy")
    parser.add_argument("--manuscript", type=Path, default=ROOT / "research/manuscript.md")
    parser.add_argument("--results", type=Path, default=ROOT / "research/results")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--repository", type=Path, default=ROOT, help="Trusted structural evidence checkout"
    )
    parser.add_argument(
        "--allow-missing-selectivity",
        action="store_true",
        help="Legacy results only: explicitly annotate absent paired-target results",
    )
    parser.add_argument(
        "--structure-input",
        type=Path,
        default=ROOT / "research/structure_comparison/results/observed_proximity.json",
    )
    parser.add_argument(
        "--structure-manifest",
        type=Path,
        default=ROOT / "research/structure_comparison/results/derived/derived_manifest.json",
    )
    parser.add_argument(
        "--handoff-directory", type=Path, default=ROOT / "research/structure_comparison"
    )
    parser.add_argument(
        "--allow-missing-structure",
        action="store_true",
        help="Legacy only: annotate absent structural evidence; never ignore malformed input",
    )
    args = parser.parse_args()
    for path in build(
        args.manuscript,
        args.results,
        args.output,
        require_selectivity=not args.allow_missing_selectivity,
        structure_input=args.structure_input,
        structure_manifest=args.structure_manifest,
        require_structure=not args.allow_missing_structure,
        handoff_directory=args.handoff_directory,
        repository=args.repository,
        profile=args.profile,
    ):
        print(path)


if __name__ == "__main__":
    main()
