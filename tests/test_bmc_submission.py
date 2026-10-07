from __future__ import annotations

import copy
import gzip
import hashlib
import io
import json
import os
import subprocess
import sys
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from build_bmc_submission import (
    CONFIRMATIONS,
    MAX_ADDITIONAL_BYTES,
    PACKAGE,
    STATEMENTS,
    article_checks,
    build,
    declarations,
    format_journal_docx,
    release_gaps,
    source_archives,
    zip_payload,
)
from build_research_documents import ROOT
from docx import Document


def test_article_limits_and_required_baselines() -> None:
    text = (ROOT / PACKAGE / "manuscript.md").read_text()
    checks = article_checks(text)
    assert checks["abstract_words_including_subheadings"] <= 200
    assert checks["body_words_including_headings_titles_legends_excluding_cells"] <= 2000
    assert checks["main_tables"] == 3 and checks["keywords"] == 6
    for name in ("Constant_0_5", "Train_prevalence", "tpsa_only_lr", "hba_only_lr", "Property_LR"):
        assert name in text
    assert "3.756 Å" in text and "3.251 Å" in text
    assert "0.600" in text and "No finite bound" in text
    assert "zero-occupancy" in text and "0.78" in text
    assert "## Declarations" not in text


@pytest.mark.parametrize(
    "mutation",
    ["abstract", "body", "sections", "objective", "citation", "title", "legend", "display"],
)
def test_format_violations_fail(mutation: str) -> None:
    text = (ROOT / PACKAGE / "manuscript.md").read_text()
    if mutation == "abstract":
        text = text.replace("### Objective\n", "### Objective\n" + "word " * 201)
    elif mutation == "body":
        text = text.replace("## Introduction\n", "## Introduction\n" + "word " * 2001)
    elif mutation == "sections":
        text += "\n## Abstract\n"
    elif mutation == "objective":
        text = text.replace("### Objective\n", "### Background\n")
    elif mutation == "citation":
        text = text.replace("### Objective\n", "### Objective\n[1]")
    elif mutation == "title":
        text = text.replace("Table 1. Deposited", "Table 1. " + "word " * 16 + "Deposited")
    elif mutation == "legend":
        text = text.replace("Site identifiers are", "word " * 301 + "Site identifiers are")
    else:
        text += "\nFigure 1. Extra display\n"
    with pytest.raises(ValueError):
        article_checks(text)


def test_author_release_is_blocked_without_facts(tmp_path: Path) -> None:
    record = json.loads((ROOT / PACKAGE / "author_confirmation.json").read_text())
    assert set(STATEMENTS).issubset(release_gaps(record))
    assert set(CONFIRMATIONS).issubset(release_gaps(record))
    with pytest.raises(ValueError, match="Author release blocked"):
        build(ROOT, tmp_path / "release", require_author_confirmation=True)
    assert not (tmp_path / "release").exists()


def approved_fixture() -> dict[str, Any]:
    return {
        "status": "author_confirmed",
        "statements": {name: f"Test fixture statement for {name}" for name in STATEMENTS},
        "confirmations": dict.fromkeys(CONFIRMATIONS, True),
    }


@pytest.mark.parametrize("bad", [None, "", "   ", "[MISSING: affiliation]", "TBD", 1])
def test_missing_and_placeholder_statements_refused(bad: Any) -> None:
    record = approved_fixture()
    record["statements"]["funding"] = bad
    assert "funding" in release_gaps(record)


@pytest.mark.parametrize("bad", [False, None, "true", 1])
def test_confirmations_must_be_explicit_booleans(bad: Any) -> None:
    record = approved_fixture()
    record["confirmations"]["contributor_history_resolved"] = bad
    assert "contributor_history_resolved" in release_gaps(record)


def test_author_statements_are_not_invented() -> None:
    record = approved_fixture()
    original = copy.deepcopy(record)
    assert release_gaps(record) == []
    text = declarations(record, {"Additional_file_2.zip": "a" * 64})
    for key in (
        "funding",
        "competing_interests",
        "author_contributions",
        "ai_assistance_disclosure",
    ):
        assert record["statements"][key] in text
    assert "Not applicable" not in text and "no funding" not in text
    assert record == original


def test_docx_is_editable_double_spaced_and_numbered(tmp_path: Path) -> None:
    path = tmp_path / "note.docx"
    doc = Document()
    doc.add_paragraph("Text")
    doc.add_table(rows=2, cols=3)
    doc.save(str(path))
    format_journal_docx(path)
    with ZipFile(path) as archive:
        xml = ElementTree.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        spacing = xml.findall(".//w:spacing", ns)
        assert spacing and all(n.attrib[f"{{{ns['w']}}}line"] == "480" for n in spacing)
        assert xml.find(".//w:lnNumType", ns) is not None
        assert b'"PAGE"' in archive.read("word/footer1.xml")
        assert b'w:type="page"' not in archive.read("word/document.xml")


def test_archives_are_deterministic_and_reject_traversal() -> None:
    files = {"a": b"first", "b": b"second"}
    assert zip_payload(files) == zip_payload(dict(reversed(list(files.items()))))
    with pytest.raises(ValueError, match="Unsafe"):
        zip_payload({"../outside": b"no"})


def test_official_source_captures_match_hashes() -> None:
    folder = ROOT / PACKAGE / "sources"
    manifest = json.loads((folder / "manifest.json").read_text())
    assert len(manifest) == 10
    for name, record in manifest.items():
        assert hashlib.sha256((folder / name).read_bytes()).hexdigest() == record["sha256"]
        assert record["url"].startswith("https://")
        assert (
            hashlib.sha256(gzip.decompress((folder / name).read_bytes())).hexdigest()
            == (record["decoded_sha256"])
        )


def test_source_archive_without_git_and_missing_map_part(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "source"
    for name in ("scripts/build_bmc_submission.py", "tests/test_bmc_submission.py", "a.ccp4.gz"):
        path = root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"fixture")
    inventory = {
        str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in root.rglob("*")
        if p.is_file()
    }
    import submission_archive

    allowlist = root / submission_archive.ALLOWLIST
    allowlist.parent.mkdir(parents=True)
    names = [*inventory, str(submission_archive.ALLOWLIST)]
    allowlist.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "files": {name: {"sha256": None, "rights_class": "fixture"} for name in names},
                "generated_files": [],
            }
        )
    )
    monkeypatch.setattr(
        submission_archive, "PUBLIC_LOCK_SHA256", hashlib.sha256(allowlist.read_bytes()).hexdigest()
    )
    inventory[str(submission_archive.ALLOWLIST)] = hashlib.sha256(
        allowlist.read_bytes()
    ).hexdigest()
    (root / "SHA256SUMS.json").write_text(json.dumps(inventory))
    generated = tmp_path / "generated"
    generated.mkdir()
    first, second = source_archives(root, generated)
    with ZipFile(io.BytesIO(first)) as archive:
        assert "source/scripts/build_bmc_submission.py" in archive.namelist()
        assert "source/a.ccp4.gz" not in archive.namelist()
    with ZipFile(io.BytesIO(second)) as archive:
        assert "source/a.ccp4.gz" in archive.namelist()
    (root / "a.ccp4.gz").write_bytes(b"changed")
    with pytest.raises(ValueError, match="inventory mismatch"):
        source_archives(root, generated)


def test_nonempty_output_refused(tmp_path: Path) -> None:
    (tmp_path / "sentinel").write_text("keep")
    with pytest.raises(ValueError, match="new or empty"):
        build(ROOT, tmp_path)
    assert (tmp_path / "sentinel").read_text() == "keep"


def test_full_package_build(tmp_path: Path) -> None:
    output = tmp_path / "note"
    manifest = json.loads(build(ROOT, output).read_text())
    assert manifest["submission_ready"] is False and manifest["submitted"] is False
    assert manifest["author_owned_gaps"]
    assert not manifest["biological_validation"]
    for name, record in manifest["files"].items():
        payload = (output / name).read_bytes()
        assert hashlib.sha256(payload).hexdigest() == record["sha256"]
        assert len(payload) == record["bytes"]
    for name in ("Additional_file_1.pdf", "Additional_file_2.zip", "Additional_file_3.zip"):
        assert (output / name).stat().st_size <= MAX_ADDITIONAL_BYTES
    with ZipFile(output / "BMC_research_note.docx") as archive:
        xml = archive.read("word/document.xml")
        assert b"AUTHOR-REVIEW DRAFT" in xml
        assert b"lnNumType" in xml
    extracted = tmp_path / "extracted"
    for name in ("Additional_file_2.zip", "Additional_file_3.zip", "Private_audit_provenance.zip"):
        with ZipFile(output / name) as archive:
            archive.extractall(extracted)
    source = extracted / "source"
    inventory = json.loads((source / "SHA256SUMS.json").read_text())
    for name, digest in inventory.items():
        assert hashlib.sha256((source / name).read_bytes()).hexdigest() == digest
    assert not (source / PACKAGE / "reviewer_candidates.md").exists()
    assert not (source / PACKAGE / "author_confirmation.json").exists()
    assert not (source / PACKAGE / "sources").exists()
    supplement = (output / "Additional_file_1.md").read_text()
    assert "## Extended structural and endpoint methods" in supplement
    assert "## Supplementary references" in supplement
    assert "The journal-specific guide was inaccessible" not in supplement
    doc = Document(str(output / "BMC_research_note.docx"))
    assert doc.tables[2].columns[0].width is not None
    assert doc.tables[2].columns[0].width.pt == 155
    assert len(Document(str(output / "Additional_file_1.docx")).inline_shapes) == 10
    rebuilt = tmp_path / "rebuilt_without_git"
    subprocess.run(
        [
            sys.executable,
            str(source / "scripts/build_bmc_submission.py"),
            "--repository",
            str(source),
            "--output",
            str(rebuilt),
        ],
        check=True,
        capture_output=True,
        text=True,
        env={**os.environ, "PYTHONPATH": str(source / "scripts/scripts")},
    )
    second_manifest = json.loads((rebuilt / "submission-manifest.json").read_text())
    assert second_manifest["article_checks"] == manifest["article_checks"]
    assert (rebuilt / "BMC_research_note.md").read_bytes() == (
        output / "BMC_research_note.md"
    ).read_bytes()
    candidate = source / PACKAGE / "manuscript.md"
    candidate.write_text(candidate.read_text().replace("3.756", "9.999", 1))
    with pytest.raises(ValueError, match="drift"):
        build(source, tmp_path / "mutated-strict", require_author_confirmation=True)
    assert not (tmp_path / "mutated-strict").exists()


@pytest.mark.parametrize(
    "bad",
    [
        "\u200b",
        "TODO",
        "[ MISSING ]",
        "ＴＢＤ",
        "T\u200bBD",
        r"\g<0>",
        "Text with TODO inside",
        "\u2800",
        "\u034f",
        "\u3164",
        "Text\n## References",
        "Text\u2028| New | Table |",
    ],
)
def test_adversarial_author_statements_fail_closed(bad: str) -> None:
    record = approved_fixture()
    record["statements"]["funding"] = bad
    assert "funding" in release_gaps(record)


@pytest.mark.parametrize(
    "payload",
    [
        "[]",
        "null",
        '"text"',
        '{"status":"author_confirmed","status":"author_answers_required"}',
        '{"statements":{"funding":"a","funding":"b"}}',
    ],
)
def test_malformed_author_json_refused(tmp_path: Path, payload: str) -> None:
    from author_release import read_record

    path = tmp_path / "record.json"
    path.write_text(payload)
    with pytest.raises(ValueError):
        release_gaps(read_record(path))


@pytest.mark.parametrize(
    "name",
    [
        "",
        "C:/outside",
        "a\\..\\outside",
        "./a",
        "a//b",
        "/absolute",
        "../escape",
        "CON",
        "dir/NUL.txt",
        "a.",
        "a ",
        "a\x00b",
        "a?b",
    ],
)
def test_portable_archive_path_refusal(name: str) -> None:
    with pytest.raises(ValueError, match="Unsafe"):
        zip_payload({name: b"no"})


def test_literal_author_text_rendering(tmp_path: Path) -> None:
    from build_research_documents import read_blocks, render_document

    text = "Actual *asterisks* and `backticks`, <angle brackets> & ampersand."
    _, path = render_document(read_blocks(text), tmp_path, [], literal_paragraphs=frozenset({text}))
    assert text in [p.text for p in Document(str(path)).paragraphs]


@pytest.mark.parametrize(
    "before,after",
    [
        ("3.756", "9.999"),
        ("0.600", "9.999"),
        ("0.9960", "0.5000"),
        ("All four", "All five"),
        ("46 rows", "47 rows"),
        ("20 versus 60", "20 versus 20"),
    ],
)
def test_prose_drift_refused_outside_tables(before: str, after: str) -> None:
    from bmc_prose import validate_prose

    inputs = [
        json.loads((ROOT / name).read_text())
        for name in (
            "research/structure_comparison/results/observed_proximity.json",
            "research/results/selectivity.json",
            "research/structure_comparison/model_support/results/model_support.json",
            "research/results/controls.json",
        )
    ]
    text = (ROOT / PACKAGE / "manuscript.md").read_text()
    validate_prose(ROOT, text, inputs)
    assert before in text
    with pytest.raises(ValueError, match="drift"):
        validate_prose(ROOT, text.replace(before, after, 1), inputs)


def test_public_inventory_excludes_private_material() -> None:
    from submission_archive import distribution_gaps, public_inventory

    names = public_inventory(ROOT)["files"]
    assert not any(
        "reviewer" in n or "author_confirmation" in n or "submission/sources" in n or "AI-R" in n
        for n in names
    )
    assert distribution_gaps(ROOT)
    assert (
        len(release_gaps(json.loads((ROOT / PACKAGE / "author_confirmation.json").read_text())))
        == 19
    )


def test_final_extra_table_anywhere_refused() -> None:
    text = (ROOT / PACKAGE / "manuscript.md").read_text()
    with pytest.raises(ValueError):
        article_checks(text + "\n| Extra | Table |\n| --- | --- |\n| a | b |\n")


def test_archive_refuses_linked_sources(tmp_path: Path) -> None:
    import os

    from submission_archive import source_bytes

    original = tmp_path / "original"
    original.write_bytes(b"data")
    (tmp_path / "symlink").symlink_to(original)
    with pytest.raises(ValueError, match="Symlink"):
        source_bytes(tmp_path, "symlink")
    (tmp_path / "dirlink").symlink_to(tmp_path, target_is_directory=True)
    with pytest.raises(ValueError, match="Symlink"):
        source_bytes(tmp_path, "dirlink/original")
    os.link(original, tmp_path / "hardlink")
    with pytest.raises(ValueError, match="single-link"):
        source_bytes(tmp_path, "hardlink")


def test_author_schema_rejects_unknown_and_missing_keys() -> None:
    record = approved_fixture()
    record["unknown"] = True
    with pytest.raises(ValueError, match="schema"):
        release_gaps(record)
    record = approved_fixture()
    del record["statements"]["funding"]
    with pytest.raises(ValueError, match="fields"):
        release_gaps(record)


@pytest.mark.parametrize("kind", ["missing", "duplicate", "incomplete", "obfuscated"])
def test_external_author_record_never_falls_back(tmp_path: Path, kind: str) -> None:
    path = tmp_path / "answers.json"
    record = approved_fixture()
    if kind == "missing":
        with pytest.raises(FileNotFoundError):
            build(ROOT, tmp_path / "out", author_record=path, require_author_confirmation=True)
    elif kind == "duplicate":
        path.write_text('{"status":"author_confirmed","status":"author_answers_required"}')
        with pytest.raises(ValueError, match="Duplicate JSON key"):
            build(ROOT, tmp_path / "out", author_record=path)
    else:
        record["statements"]["funding"] = None if kind == "incomplete" else "T\u034fB\u034fD"
        path.write_text(json.dumps(record))
        with pytest.raises(ValueError, match="Author release blocked.*funding"):
            build(ROOT, tmp_path / "out", author_record=path, require_author_confirmation=True)
    assert not (tmp_path / "out").exists()


def test_completed_author_path_without_unpinning_provenance(tmp_path: Path) -> None:
    import shutil

    from submission_archive import distribution_gaps

    # Synthetic declarations/permissions exercise software, never approve the actual paper.
    checkout = tmp_path / "checkout"
    subprocess.run(
        ["git", "clone", "--shared", str(ROOT), str(checkout)],
        check=True,
        capture_output=True,
    )
    tracked = subprocess.check_output(["git", "-C", str(ROOT), "ls-files", "-z"]).decode()
    for name in filter(None, tracked.split("\0")):
        (checkout / name).parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT / name, checkout / name)
    record_path = tmp_path / "completed-record.json"
    record = approved_fixture()
    record["statements"]["final_byline_and_addresses"] = "Synthetic test author, José *literal*."
    record_bytes = json.dumps(record, ensure_ascii=False).encode()
    record_path.write_bytes(record_bytes)
    rights_path = checkout / PACKAGE / "rights_review.json"
    rights = json.loads(rights_path.read_bytes())
    rights["project_code_license"] = "Synthetic test license, not a rights decision."
    rights["anonymous_access_checked"] = True
    for row in rights["classes"].values():
        for key in row:
            row[key] = True if key == "author_cleared" else "Synthetic test assertion only."
    for row in rights["items"].values():
        row["author_cleared"] = True
        row["decision_and_source"] = "Synthetic test assertion only."
    rights_path.write_text(json.dumps(rights))
    assert distribution_gaps(checkout) == []
    frozen_record = (checkout / PACKAGE / "author_confirmation.json").read_bytes()
    output = tmp_path / "confirmed"
    manifest = json.loads(
        build(
            checkout, output, author_record=record_path, require_author_confirmation=True
        ).read_bytes()
    )
    assert manifest["author_record_sha256"] == hashlib.sha256(record_bytes).hexdigest()
    assert manifest["external_author_record"] and manifest["submission_ready"]
    assert not manifest["submitted"] and not manifest["biological_validation"]
    assert (output / "author_confirmation.json").read_bytes() == record_bytes
    assert (checkout / PACKAGE / "author_confirmation.json").read_bytes() == frozen_record
    article = (output / "BMC_research_note.md").read_text()
    assert "AUTHOR-REVIEW DRAFT" not in article and "## Declarations" in article
    for key in ("final_byline_and_addresses", "funding", "ai_assistance_disclosure"):
        assert record["statements"][key] in article
    doc = Document(str(output / "BMC_research_note.docx"))
    assert record["statements"]["final_byline_and_addresses"] in [p.text for p in doc.paragraphs]
    with ZipFile(output / "Private_audit_provenance.zip") as archive:
        assert archive.read(f"private/source/{PACKAGE}/author_confirmation.json") == frozen_record
    extracted = tmp_path / "extracted"
    for name in ("Additional_file_2.zip", "Additional_file_3.zip", "Private_audit_provenance.zip"):
        with ZipFile(output / name) as archive:
            if name != "Private_audit_provenance.zip":
                assert not any(
                    "author_confirmation" in n or "completed-record" in n
                    for n in archive.namelist()
                )
            archive.extractall(extracted)
    source = extracted / "source"
    rebuilt = tmp_path / "confirmed-without-git"
    subprocess.run(
        [
            sys.executable,
            str(source / "scripts/build_bmc_submission.py"),
            "--repository",
            str(source),
            "--output",
            str(rebuilt),
            "--author-record",
            str(record_path),
            "--require-author-confirmation",
        ],
        check=True,
        capture_output=True,
        env={**os.environ, "PYTHONPATH": str(source / "scripts/scripts")},
    )
    normalized = []
    for package in (output, rebuilt):
        built_manifest = json.loads((package / "submission-manifest.json").read_bytes())
        prose = (package / "BMC_research_note.md").read_text()
        for name in ("Additional_file_2.zip", "Additional_file_3.zip"):
            digest = hashlib.sha256((package / name).read_bytes()).hexdigest()
            assert digest == built_manifest["files"][name]["sha256"]
            assert prose.count(digest) == 1
            prose = prose.replace(digest, f"{name} verified digest")
        prefix = "Code access and requirements: "
        start = prose.index(prefix) + len(prefix)
        access, length = json.JSONDecoder().raw_decode(prose[start:])
        assert access == built_manifest["code_access"]
        normalized.append(
            prose[:start] + "Verified build-specific source provenance" + prose[start + length :]
        )
    # Archive bytes and Git availability differ; scientific prose and author text must not.
    assert normalized[0] == normalized[1]
    assert (rebuilt / "author_confirmation.json").read_bytes() == record_bytes
    # The override must not excuse tampering with the archived original author record.
    (checkout / PACKAGE / "author_confirmation.json").write_text(json.dumps(record))
    with pytest.raises(ValueError, match="Private provenance changed:.*author_confirmation"):
        build(
            checkout,
            tmp_path / "tampered",
            author_record=record_path,
            require_author_confirmation=True,
        )
    assert not (tmp_path / "tampered").exists()
